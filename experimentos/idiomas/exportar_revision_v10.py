"""
Exporta los targets v10 a tres xlsx para leerlos (uno por celda):
attributes/v10/revision_v10_<celda>.xlsx

Hoja `resumen`: por categoria, filas, gate, accuracy, idioma y cuantas se corrigieron a mano.
Hoja `respuestas`: una fila por pregunta con la pregunta en ingles y en el idioma de la
celda, la respuesta canonica, el target final y, si se corrigio a mano, el motivo y el
output original del modelo. Las filas corregidas van resaltadas.

El output original sale del commit donde se generaron los targets (--ref, c0b1e57 =
"targets full"), porque correcciones_v10.json solo guarda su sha1.

    python3 exportar_revision_v10.py [--dir attributes/v10] [--ref c0b1e57]
"""

import argparse
import io
import json
import os
import subprocess

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SPLIT = 0.84
CORRECCIONES = "data/banco_v8/correcciones_v10.json"
RESALTE = PatternFill("solid", fgColor="FFF2CC")


def originales(ref, path):
    rel = os.path.relpath(os.path.abspath(path), _raiz_git())
    try:
        txt = subprocess.run(["git", "show", f"{ref}:{rel}"], check=True, capture_output=True, text=True).stdout
    except subprocess.CalledProcessError:
        print(f"aviso: no encuentro {rel} en {ref}; la columna output_original queda vacia")
        return {}
    df = pd.read_csv(io.StringIO(txt), sep=";", keep_default_na=False, dtype=str)
    return dict(zip(df["prompt"], df["output"]))


def _raiz_git():
    return subprocess.run(["git", "rev-parse", "--show-toplevel"], check=True,
                          capture_output=True, text=True).stdout.strip()


def hoja(ws, filas, anchos):
    for r in filas:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dir", default="attributes/v10")
    ap.add_argument("--ref", default="c0b1e57")
    args = ap.parse_args()
    with open(CORRECCIONES, encoding="utf-8") as f:
        corr = json.load(f)

    for celda in ("fr", "es", "de"):
        path = os.path.join(args.dir, f"targets_v8_{celda}.csv")
        df = pd.read_csv(path, sep=";", keep_default_na=False, dtype=str)
        orig = originales(args.ref, path)
        c = corr.get(celda, {})
        corte = int(len(df) * SPLIT)
        es_true = lambda s: s.str.lower() == "true"  # noqa: E731

        wb = Workbook()
        ws = wb.active
        ws.title = "resumen"
        filas = [["categoria", "filas", "gate", "accuracy", f"idioma == {celda}", "corregidas a mano"]]
        for cat, g in list(df.groupby("categoria", sort=False)) + [("TOTAL", df)]:
            filas.append([cat, len(g), int(es_true(g["passed_gate"]).sum()), int(es_true(g["ref_answer_correct"]).sum()),
                          int((g["ref_language"] == celda).sum()), int(g["prompt"].isin(c).sum())])
        hoja(ws, filas, {"A": 26, "B": 8, "C": 8, "D": 10, "E": 12, "F": 18})

        ws = wb.create_sheet("respuestas")
        filas = [["fila", "split", "categoria", "prompt", f"prompt_{celda}", "answer", "output",
                  "corregida", "motivo", "output_original", "gate", "correcta", "idioma"]]
        for i, r in df.iterrows():
            cc = c.get(r["prompt"])
            filas.append([i, "train" if i < corte else "held-out", r["categoria"], r["prompt"], r[f"prompt_{celda}"],
                          r["answer"], r["output"], "si" if cc else "", cc["motivo"] if cc else "",
                          orig.get(r["prompt"], "") if cc else "", r["passed_gate"], r["ref_answer_correct"],
                          r["ref_language"]])
        hoja(ws, filas, {"A": 6, "B": 9, "C": 20, "D": 36, "E": 36, "F": 16, "G": 70, "H": 10,
                         "I": 40, "J": 70, "K": 7, "L": 9, "M": 8})
        envolver = Alignment(wrap_text=True, vertical="top")
        for fila in ws.iter_rows(min_row=2):
            for cel in fila:
                cel.alignment = envolver
            if fila[7].value == "si":
                for cel in fila:
                    cel.fill = RESALTE
        ws.auto_filter.ref = ws.dimensions

        out = os.path.join(args.dir, f"revision_v10_{celda}.xlsx")
        wb.save(out)
        print(f"[{celda}] {out}: {len(df)} filas, {int(df['prompt'].isin(c).sum())} corregidas")


if __name__ == "__main__":
    main()
